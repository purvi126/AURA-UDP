"""Build the editable AURA-UDP Review 2 progress report from checked-in evidence."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BLUE = "173B59"
PALE = "EAF1F6"
GRAY = "D9D9D9"


def shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def border(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "5")
        element.set(qn("w:color"), GRAY)
        borders.append(element)


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, header in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.width = Inches(widths[i])
        cell.text = header
        shade(cell, BLUE)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
        border(cell)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = Inches(widths[i])
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri % 2:
                shade(cells[i], "F4F7F9")
            border(cells[i])
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.12
                for run in p.runs:
                    run.font.size = Pt(9)
            margins = cell._tc.get_or_add_tcPr()
            mar = OxmlElement("w:tcMar")
            for side in ("top", "start", "bottom", "end"):
                e = OxmlElement(f"w:{side}")
                e.set(qn("w:w"), "90")
                e.set(qn("w:type"), "dxa")
                mar.append(e)
            margins.append(mar)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def add_text(doc, text, *, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)
    return p


def page_break(doc):
    doc.add_page_break()


def build(output: Path, chart_png: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(0.76)
    sec.right_margin = Inches(0.76)
    sec.header_distance = Inches(0.3)
    sec.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(25, 32, 38)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.17
    for name, size, before, after in [
        ("Title", 20, 0, 13),
        ("Heading 1", 13, 13, 6),
        ("Heading 2", 11, 10, 4),
    ]:
        s = styles[name]
        s.font.name = "Aptos"
        s.font.size = Pt(size)
        s.font.bold = name != "Title"
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True
    styles["Title"].font.bold = True
    # Word's built-in Title style can carry an accent border in some themes.
    title_ppr = styles["Title"].element.get_or_add_pPr()
    for existing in list(title_ppr):
        if existing.tag == qn("w:pBdr"):
            title_ppr.remove(existing)
    styles["Caption"].font.name = "Aptos"
    styles["Caption"].font.size = Pt(8.5)
    styles["Caption"].font.color.rgb = RGBColor(70, 79, 87)
    styles["Caption"].paragraph_format.space_after = Pt(8)
    styles["List Bullet"].paragraph_format.space_after = Pt(4)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("15 October 2026  •  Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

    doc.add_paragraph("AURA-UDP implementation progress report", style="Title")
    p = doc.add_paragraph()
    p.add_run("Review 2 milestone  •  Approximately 70% complete").bold = True
    p.add_run("\nDeadline-aware reliable transport for critical campus IoT messages")
    add_text(doc, "This report explains the project for a reader who has not seen the code. It records what is implemented and verified, what the NS2 experiments found, what Cisco Packet Tracer demonstrates, and what remains for the final review. The 70% figure is a planning estimate of overall project readiness, not a measured software metric.")
    doc.add_heading("The project in one minute", level=1)
    add_text(doc, "A smart campus sends routine readings and urgent alerts to a monitoring room. A lost or late routine reading may be tolerable; a smoke alert is useful only if it arrives before its deadline. AURA-UDP asks whether a sender can improve timely alert delivery while controlling the cost of retransmissions.")
    add_text(doc, "Two tools show different parts of that question. NS2 runs the actual transport experiment and produces timing, loss, and retry measurements. Cisco Packet Tracer is an editable campus network model that shows routers, HTTP connectivity, an IoT gateway, and registered sensor devices. The tools do not exchange live packets.")
    doc.add_heading("What the 70% milestone covers", level=1)
    table(doc, ["Workstream", "Implemented and evidenced", "Status"], [
        ("Transport in NS2", "Baseline, Selective Repeat, and AURA sender/receiver behavior; ACKs, deadlines, retry and congestion-window decisions.", "Built"),
        ("Evaluation", "Three scenarios, paired seed sweeps, policy ablations, fresh-seed checks, metrics, charts, and a local dashboard.", "Built"),
        ("Campus in Packet Tracer", "Two routers, monitoring server, native temperature and smoke devices, gateway registration, routed HTTP check.", "Built"),
        ("Final research claim and demo", "Wider workloads, stronger fairness/congestion tests, improved retry estimate, and final presentation evidence.", "Remaining"),
    ], [1.45, 4.65, 0.83])
    doc.add_heading("How the pieces fit", level=1)
    add_text(doc, "Campus alert idea  →  NS2 sender and lossy link  →  event logs and metrics  →  comparison charts and dashboard")
    add_text(doc, "Campus layout idea  →  Packet Tracer routers, server, gateway, and sensors  →  connectivity and registration screenshots")

    page_break(doc)
    doc.add_heading("What happens inside NS2", level=1)
    add_text(doc, "The script ns2/run.tcl builds a sender, router, monitoring receiver, and competing CBR sender. Its 128 Kb/s bottleneck has a short queue and configurable packet loss. Each run creates 30 messages; every fourth is critical and has an absolute deadline. Data and ACK paths have separate loss streams.")
    table(doc, ["Mode", "Sender behavior", "Purpose"], [
        ("Baseline UDP", "Sends a message once; no ACK or retry.", "Reference for low overhead."),
        ("Fixed Selective Repeat", "Uses a four-message window, ACKs and individual retries with a fixed timer.", "Reference reliable method."),
        ("AURA", "Orders urgent alerts by deadline, gives routine traffic regular turns, adapts timer and small window, and checks whether a retry can still arrive in time.", "Tests the deadline-aware policy."),
    ], [1.48, 4.35, 1.1])
    doc.add_heading("The decision being tested", level=2)
    add_text(doc, "For a critical message, the sender compares estimated forward delivery time with time remaining until its absolute deadline. If the estimated trip is longer, AURA abandons that retransmission. The first send is still attempted while the alert is live. This can save traffic, but a rough estimate may reject a retry that would have succeeded.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Retry when estimated delivery time ≤ deadline − current time")
    run.bold = True
    doc.add_heading("What is measured", level=2)
    add_bullet(doc, "Deadline delivery ratio (DDR): critical messages first delivered on time divided by all generated critical messages, including lost and abandoned ones.")
    add_bullet(doc, "Packet delivery ratio (PDR): unique messages delivered eventually divided by all generated messages. A late alert can count in PDR but not DDR.")
    add_bullet(doc, "Also recorded: delay for delivered messages, extra sends per generated message, routine delivery, and received background traffic.")
    add_text(doc, "NS2 writes events.csv, control.csv, background.csv, a packet trace, and a NAM trace. Python analysis deduplicates message IDs and produces summary CSVs, plots, and the local dashboard. The dashboard reads saved results; it does not run NS2.")
    doc.add_heading("Experiment design", level=2)
    add_text(doc, "The development set uses 30 seed indices across three conditions and three transport modes (270 runs), plus four AURA variants (360 runs). A fresh 30-seed set repeats both sweeps. Total recorded evaluation: 1,260 NS2 runs. Within a scenario, modes use the same topology, generation schedule, and seed index; changed retry behavior means packet-by-packet loss draws are not identical.")

    page_break(doc)
    doc.add_heading("NS2 evidence and measured findings", level=1)
    add_text(doc, "The figure below is generated from the checked-in NS2 run summaries. It compares delivery and retry behavior across moderate loss, tight deadlines, and a higher message rate. It is experimental output, not Packet Tracer traffic.")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(chart_png), width=Inches(6.35))
    add_caption(doc, "Figure 1  NS2 30-seed comparison from experiments/plots/multiseed.svg")
    table(doc, ["Scenario", "Baseline DDR", "Fixed SR DDR", "AURA DDR"], [
        ("Moderate loss", "81.7%", "83.3%", "93.3%"),
        ("Tight deadline", "12.9%", "10.0%", "10.4%"),
        ("High message rate", "86.3%", "73.3%", "89.2%"),
    ], [2.03, 1.6, 1.6, 1.7])
    add_text(doc, "AURA delivered more critical messages on time than fixed SR in the moderate and high-rate conditions. The tight-deadline condition remained difficult for every mode, and baseline exceeded AURA there. The feasibility check reduced high-rate retry cost in the development set, but did not show a DDR gain; on fresh seeds, full AURA and the version without this check both reached 90.4% DDR. These are descriptive results for three short synthetic workloads.")

    page_break(doc)
    doc.add_heading("Cisco Packet Tracer implementation", level=1)
    add_text(doc, "The editable packet_tracer/topology/AURA-UDP-campus.pkt models the campus side of the project. PC0 on the sensor LAN reaches Server0 at 10.10.30.10 through Router0 and Router1. Separately, a native Temperature Monitor and Smoke Detector register with HomeGateway0; Tablet PC0 displays them at 192.168.25.1.")
    add_text(doc, "The original captures below are kept at a modest size because Packet Tracer exported them at only 513 pixels wide. Their captions state the tested observation so the small text is not the only evidence.")
    t = doc.add_table(rows=1, cols=2)
    t.autofit = False
    captures = [
        (ROOT / "packet_tracer/documentation/media/routed-http-pc0.jpg", "PC0 browser: routed HTTP page served by Server0."),
        (ROOT / "packet_tracer/documentation/media/gateway-devices.jpg", "Tablet gateway view: smoke and temperature devices registered."),
    ]
    for cell, (path, caption) in zip(t.rows[0].cells, captures):
        cell.width = Inches(3.42)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(3.17))
        cp = cell.add_paragraph()
        cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cp.add_run(caption)
        r.font.size = Pt(8.5)
        r.italic = True
    add_caption(doc, "Figures 2 and 3  Original Cisco Packet Tracer screenshots")
    doc.add_heading("What these checks establish", level=2)
    add_bullet(doc, "The saved PC0 browser loaded the custom server page over the two-router path.")
    add_bullet(doc, "The gateway device list showed both native IoT devices and their displayed fields.")
    add_text(doc, "Packet Tracer does not execute AURA or feed the NS2 simulation. A triggered smoke alarm transition has not been verified. The gateway and wireless settings are a lab demonstration, not a production security design.")

    page_break(doc)
    doc.add_heading("Verification and remaining work", level=1)
    add_text(doc, "The repository records 19 passing Python and NS2 checks, repeatable sweep outputs, and a GitHub Actions workflow that runs tests and small NS2 smoke sweeps. The saved Packet Tracer checks cover routed HTTP, pings, and native IoT registration. Versioned Git commits and releases provide rollback points.")
    doc.add_heading("Remaining 30% for the final review", level=2)
    add_bullet(doc, "Test longer and more varied workloads, more competing traffic, and a wider range of deadlines and loss conditions.")
    add_bullet(doc, "Improve the forward-time estimate with queueing uncertainty and evaluate whether the deadline decision improves DDR without starving routine traffic.")
    add_bullet(doc, "Strengthen congestion and fairness evaluation; the current bounded window is a lab heuristic, not an Internet-grade controller.")
    add_bullet(doc, "Capture a verified smoke-event transition in Packet Tracer and prepare a concise live demonstration and final report.")
    doc.add_heading("How to reproduce the evidence", level=2)
    add_text(doc, "Open packet_tracer/topology/AURA-UDP-campus.pkt in Packet Tracer 9.0. On PC0, browse to http://10.10.30.10. On Tablet PC0, browse to http://192.168.25.1 and inspect the gateway device list.")
    add_text(doc, "In Ubuntu or WSL with NS2 2.35 installed, run python3 scripts/multiseed.py and python3 scripts/ablation.py from the repository root. Open experiments/plots/multiseed.svg and docs/RESEARCH_EVALUATION.md for the measured comparison. To inspect the local dashboard, start python3 -m http.server 8000 --bind 127.0.0.1 from the repository root, then open http://localhost:8000/python/dashboard/.")
    doc.add_heading("Source files in the repository", level=2)
    table(doc, ["File or folder", "What it contains"], [
        ("ns2/run.tcl", "Topology, transport behavior, deadlines, ACKs, and logs."),
        ("scripts/ and python/analysis/", "Experiment runners, metrics, and plots."),
        ("experiments/results/", "Per-run and summary data for development and fresh seeds."),
        ("packet_tracer/", "Editable campus topology, web page, build guide, and captures."),
        ("docs/PROTOCOL.md and docs/VALIDATION.md", "Decision rules and verification record."),
    ], [2.65, 4.28])
    add_text(doc, "Review 2 conclusion: the core simulation, comparison pipeline, campus model, and presentation evidence are implemented. The final phase should focus on broader validation and a clearer demonstration of when the deadline-aware decision helps or fails.")

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument(
        "--chart-png",
        type=Path,
        default=ROOT / "docs/reviews/assets/multiseed.png",
    )
    args = parser.parse_args()
    build(args.output, args.chart_png)
